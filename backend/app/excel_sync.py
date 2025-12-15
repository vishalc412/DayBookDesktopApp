"""
Real-Time Excel Synchronization Module
Handles reading and writing daybook entries to Excel files
"""

import os
from datetime import datetime
from typing import List, Optional
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

from .models import DaybookEntry


class ExcelSyncHandler(FileSystemEventHandler):
    """Handler for Excel file changes"""

    def __init__(self, callback):
        self.callback = callback
        super().__init__()

    def on_modified(self, event):
        if not event.is_directory and event.src_path.endswith('.xlsx'):
            self.callback(event.src_path)


class ExcelSync:
    """Real-time Excel synchronization manager"""

    HEADERS = ['Entry ID', 'Date', 'Description', 'Debit', 'Credit', 'Balance', 'Category', 'Reference']

    def __init__(self, excel_file_path: str):
        self.excel_file_path = excel_file_path
        self.observer = None
        self._ensure_file_exists()

    def _ensure_file_exists(self):
        """Create Excel file if it doesn't exist"""
        if not os.path.exists(self.excel_file_path):
            self.create_new_workbook()

    def create_new_workbook(self):
        """Create a new formatted Excel workbook"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Daybook"

        # Style definitions
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Add headers
        for col, header in enumerate(self.HEADERS, start=1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # Set column widths
        column_widths = [10, 12, 30, 15, 15, 15, 15, 15]
        for col, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width

        # Freeze header row
        ws.freeze_panes = 'A2'

        wb.save(self.excel_file_path)

    def read_entries(self) -> List[DaybookEntry]:
        """Read all entries from Excel file"""
        entries = []

        try:
            wb = load_workbook(self.excel_file_path)
            ws = wb.active

            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None:  # Check if entry_id exists
                    entry = DaybookEntry(
                        entry_id=row[0],
                        date=row[1] if isinstance(row[1], datetime) else datetime.now(),
                        description=row[2] or "",
                        debit=float(row[3]) if row[3] else 0.0,
                        credit=float(row[4]) if row[4] else 0.0,
                        balance=float(row[5]) if row[5] else 0.0,
                        category=row[6] or "",
                        reference=row[7] or ""
                    )
                    entries.append(entry)

            wb.close()
        except Exception as e:
            print(f"Error reading Excel file: {e}")

        return entries

    def write_entries(self, entries: List[DaybookEntry]):
        """Write entries to Excel file"""
        try:
            wb = load_workbook(self.excel_file_path)
            ws = wb.active

            # Clear existing data (keep header)
            ws.delete_rows(2, ws.max_row)

            # Add entries
            for idx, entry in enumerate(entries, start=2):
                ws.cell(row=idx, column=1, value=entry.entry_id)
                ws.cell(row=idx, column=2, value=entry.date)
                ws.cell(row=idx, column=3, value=entry.description)
                ws.cell(row=idx, column=4, value=entry.debit)
                ws.cell(row=idx, column=5, value=entry.credit)
                ws.cell(row=idx, column=6, value=entry.balance)
                ws.cell(row=idx, column=7, value=entry.category)
                ws.cell(row=idx, column=8, value=entry.reference)

                # Format numbers
                ws.cell(row=idx, column=4).number_format = '#,##0.00'
                ws.cell(row=idx, column=5).number_format = '#,##0.00'
                ws.cell(row=idx, column=6).number_format = '#,##0.00'

            wb.save(self.excel_file_path)
            wb.close()
        except Exception as e:
            print(f"Error writing to Excel file: {e}")

    def add_entry(self, entry: DaybookEntry) -> bool:
        """Add a single entry to Excel file"""
        try:
            entries = self.read_entries()

            # Auto-generate entry_id if not provided
            if entry.entry_id is None:
                entry.entry_id = len(entries) + 1

            # Calculate balance
            if entries:
                last_balance = entries[-1].balance
                entry.balance = last_balance + entry.debit - entry.credit
            else:
                entry.balance = entry.debit - entry.credit

            entries.append(entry)
            self.write_entries(entries)
            return True
        except Exception as e:
            print(f"Error adding entry: {e}")
            return False

    def update_entry(self, entry_id: int, updated_entry: DaybookEntry) -> bool:
        """Update an existing entry"""
        try:
            entries = self.read_entries()

            for idx, entry in enumerate(entries):
                if entry.entry_id == entry_id:
                    entries[idx] = updated_entry
                    entries[idx].entry_id = entry_id

                    # Recalculate balances for all subsequent entries
                    self._recalculate_balances(entries, idx)
                    self.write_entries(entries)
                    return True

            return False
        except Exception as e:
            print(f"Error updating entry: {e}")
            return False

    def delete_entry(self, entry_id: int) -> bool:
        """Delete an entry"""
        try:
            entries = self.read_entries()

            for idx, entry in enumerate(entries):
                if entry.entry_id == entry_id:
                    del entries[idx]

                    # Recalculate balances
                    self._recalculate_balances(entries, idx)
                    self.write_entries(entries)
                    return True

            return False
        except Exception as e:
            print(f"Error deleting entry: {e}")
            return False

    def _recalculate_balances(self, entries: List[DaybookEntry], start_idx: int = 0):
        """Recalculate balances starting from a specific index"""
        for idx in range(start_idx, len(entries)):
            if idx == 0:
                entries[idx].balance = entries[idx].debit - entries[idx].credit
            else:
                prev_balance = entries[idx - 1].balance
                entries[idx].balance = prev_balance + entries[idx].debit - entries[idx].credit

    def start_watching(self, callback):
        """Start watching Excel file for changes"""
        event_handler = ExcelSyncHandler(callback)
        self.observer = Observer()
        watch_path = str(Path(self.excel_file_path).parent)
        self.observer.schedule(event_handler, watch_path, recursive=False)
        self.observer.start()

    def stop_watching(self):
        """Stop watching Excel file"""
        if self.observer:
            self.observer.stop()
            self.observer.join()
