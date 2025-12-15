"""
Excel Sync Service
Event-driven Excel synchronization
Triggers on TRANSACTION_POSTED events
"""

import os
from decimal import Decimal
from datetime import datetime
from typing import List
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.core.events import DomainEvent, EventType
from app.shared.enums import TransactionStatus


class ExcelSyncService:
    """
    Excel synchronization service
    Writes posted transactions to Excel file
    """

    HEADERS = ['Entry Number', 'Date', 'Description', 'Account', 'Debit', 'Credit', 'Reference', 'Status']

    def __init__(self, excel_file_path: str = None):
        self.excel_file_path = excel_file_path or settings.EXCEL_FILE_PATH
        self._ensure_file_exists()

    def _ensure_file_exists(self):
        """Create Excel file if it doesn't exist"""
        if not os.path.exists(self.excel_file_path):
            self.create_new_workbook()

    def create_new_workbook(self):
        """Create a new formatted Excel workbook"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Journal Entries"

        # Style definitions
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Add headers
        for col, header in enumerate(self.HEADERS, start=1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # Set column widths
        column_widths = [15, 12, 40, 30, 15, 15, 20, 10]
        for col, width in enumerate(column_widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width

        # Freeze header row
        ws.freeze_panes = 'A2'

        wb.save(self.excel_file_path)

    async def sync_transaction(self, journal_entry):
        """
        Sync a single transaction to Excel
        Called when a transaction is posted
        """
        try:
            wb = load_workbook(self.excel_file_path)
            ws = wb.active

            # Find next row
            next_row = ws.max_row + 1

            # Write transaction lines
            for line in journal_entry.lines:
                ws.cell(row=next_row, column=1, value=journal_entry.entry_number)
                ws.cell(row=next_row, column=2, value=journal_entry.date)
                ws.cell(row=next_row, column=3, value=journal_entry.description)

                # Get account name (we'll need to fetch this)
                account_name = line.account_id  # TODO: Get actual account name
                ws.cell(row=next_row, column=4, value=account_name)

                ws.cell(row=next_row, column=5, value=float(line.debit))
                ws.cell(row=next_row, column=6, value=float(line.credit))
                ws.cell(row=next_row, column=7, value=journal_entry.reference or "")
                ws.cell(row=next_row, column=8, value=journal_entry.status.value)

                # Format numbers
                ws.cell(row=next_row, column=5).number_format = '#,##0.00'
                ws.cell(row=next_row, column=6).number_format = '#,##0.00'

                next_row += 1

            wb.save(self.excel_file_path)
            wb.close()
            return True

        except Exception as e:
            print(f"Error syncing to Excel: {e}")
            return False

    async def rebuild_workbook(self, transactions: List):
        """
        Rebuild entire workbook from transactions
        Used for recovery or full sync
        """
        self.create_new_workbook()

        for transaction in transactions:
            if transaction.status == TransactionStatus.POSTED:
                await self.sync_transaction(transaction)


# Event handlers
async def on_transaction_posted(event: DomainEvent):
    """
    Event handler for TRANSACTION_POSTED
    Automatically syncs to Excel
    """
    from app.core.database import get_db_context
    from app.modules.transactions.repository import TransactionRepository

    async with get_db_context() as db:
        repo = TransactionRepository(db)
        entry = await repo.get_by_id(event.data["entry_id"])

        if entry:
            sync_service = ExcelSyncService()
            await sync_service.sync_transaction(entry)
            print(f"Synced transaction {entry.entry_number} to Excel")


async def on_transaction_reversed(event: DomainEvent):
    """
    Event handler for TRANSACTION_REVERSED
    Syncs reversal entry to Excel
    """
    from app.core.database import get_db_context
    from app.modules.transactions.repository import TransactionRepository

    async with get_db_context() as db:
        repo = TransactionRepository(db)
        reversal_entry = await repo.get_by_id(event.data["reversal_entry_id"])

        if reversal_entry:
            sync_service = ExcelSyncService()
            await sync_service.sync_transaction(reversal_entry)
            print(f"Synced reversal {reversal_entry.entry_number} to Excel")


def register_excel_sync_handlers():
    """Register Excel sync event handlers"""
    from app.core.events import event_bus

    event_bus.subscribe(EventType.TRANSACTION_POSTED, on_transaction_posted)
    event_bus.subscribe(EventType.TRANSACTION_REVERSED, on_transaction_reversed)
